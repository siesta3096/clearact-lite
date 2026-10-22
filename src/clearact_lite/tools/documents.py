import asyncio
import json
from pathlib import Path

from clearact_lite.core.contracts import ToolResult, ToolSpec
from clearact_lite.core.ports import ToolContext
from clearact_lite.storage.snapshots import workspace_path
from clearact_lite.tools.files import object_schema


def extract_document(path: Path, page: int | None, sheet: str | None) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = [page - 1] if page is not None else range(len(reader.pages))
        return "\n\n".join(f"[Page {index + 1}]\n{reader.pages[index].extract_text() or ''}" for index in pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = [p.text for p in document.paragraphs]
        parts += ["\n".join("\t".join(c.text for c in row.cells) for row in table.rows) for table in document.tables]
        return "\n".join(parts)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False)
        try:
            names = [sheet] if sheet else book.sheetnames
            return "\n\n".join(
                f"[Sheet {name}]\n"
                + "\n".join(
                    "\t".join("" if value is None else str(value) for value in row)
                    for row in book[name].iter_rows(values_only=True)
                )
                for name in names
            )
        finally:
            book.close()
    raise ValueError("document_read supports PDF, DOCX, XLSX/XLSM. Use file_read for text/CSV/JSON.")


class ReadDocument:
    spec = ToolSpec(
        "document_read",
        "Extract PDF/DOCX/XLSX text or formulas, paginated by characters. "
        "PDF page is 1-based; spreadsheets accept an optional sheet name. Scanned PDFs need OCR.",
        object_schema(
            {
                "path": {"type": "string"},
                "offset": {"type": "integer", "minimum": 0},
                "limit": {"type": "integer", "minimum": 1, "maximum": 12000},
                "page": {"type": "integer", "minimum": 1},
                "sheet": {"type": "string"},
            },
            ["path"],
        ),
    )

    async def execute(self, arguments: dict, context: ToolContext) -> ToolResult:
        path = workspace_path(context.workspace, arguments["path"])
        text = await asyncio.to_thread(extract_document, path, arguments.get("page"), arguments.get("sheet"))
        offset, limit = arguments.get("offset", 0), arguments.get("limit", 6000)
        return ToolResult(
            json.dumps(
                {
                    "path": arguments["path"],
                    "characters": len(text),
                    "offset": offset,
                    "text": text[offset : offset + limit],
                },
                ensure_ascii=False,
            )
        )
