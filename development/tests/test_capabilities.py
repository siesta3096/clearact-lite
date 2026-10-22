import asyncio
import io
import json
import sys
from pathlib import Path

import pytest
from mcp_fixture import make_server
from test_gateway import gateway, wait_idle
from verify_frontend import LocalServer

from clearact_lite.bootstrap import build_application
from clearact_lite.core.contracts import Message, ModelReply, ToolCall, Usage
from clearact_lite.integrations.config import normalize_servers
from clearact_lite.runtime.runner import ExecutionContext
from clearact_lite.settings import Settings
from clearact_lite.storage.uploads import UploadStore
from clearact_lite.tools.documents import extract_document


def stdio_config(root):
    return {
        "fixture": {
            "command": sys.executable,
            "args": [str(Path(__file__).resolve().parents[1] / "scripts/mcp_fixture.py"), "--root", str(root)],
            "env": {"FIXTURE_SECRET": "never-show-this"},
        }
    }


def context(app, branch):
    return ExecutionContext(app.settings.workspace, branch, app.store, app.snapshots, app._emit(branch))


def test_upload_original_duplicate_names_and_changed_copy(tmp_path):
    store = UploadStore(tmp_path)
    first, second = store.add("notes.txt", b"first"), store.add("notes.txt", b"second")
    assert first["path"] != second["path"]
    (tmp_path / first["path"]).write_bytes(b"edited")
    with pytest.raises(ValueError):
        store.get(first["id"])
    assert (store.root / "blobs" / first["sha256"]).read_bytes() == b"first"
    safe = store.add("../../CON.txt", b"")
    assert ".." not in safe["path"] and (tmp_path / safe["path"]).is_file()
    with pytest.raises(ValueError):
        store.get("../../escape")


def test_http_upload_linking_and_size_limit(tmp_path):
    class CaptureProvider:
        messages = []

        async def complete(self, messages, tools, on_delta):
            self.messages = messages
            return ModelReply(Message("assistant", "recorded"), "stop", Usage(10, 2))

    provider = CaptureProvider()

    async def scenario():
        async with gateway(tmp_path, provider) as (client, app):
            response = await client.post("/api/uploads", files={"file": ("中文.txt", b"proof", "text/plain")})
            assert response.status_code == 201
            info = response.json()["file"]
            run = await client.post("/api/runs", json={"prompt": "Read this", "attachment_ids": [info["id"]]})
            await wait_idle(client)
            report = (await client.get(f"/api/branches/{run.json()['branch_id']}")).json()
            linked = next(e for e in report["events"] if e["kind"] == "attachment.linked")
            assert linked["payload"]["prompt"] == "Read this" and info["path"] in provider.messages[-1].content
            app.state.service.upload_limit_bytes = 2
            assert (await client.post("/api/uploads", files={"file": ("x", b"large")})).status_code == 413
            bad = await client.post("/api/runs", json={"prompt": "x", "attachment_ids": ["../../bad"]})
            assert bad.status_code == 400

    asyncio.run(scenario())


def test_native_document_extraction_and_binary_restore(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    document = Document()
    document.add_paragraph("DOCX proof")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "table proof"
    document.save(tmp_path / "sample.docx")
    assert "table proof" in extract_document(tmp_path / "sample.docx", None, None)
    workbook = Workbook()
    workbook.active["A1"], workbook.active["B1"] = "XLSX proof", "=SUM(1,2)"
    workbook.save(tmp_path / "sample.xlsx")
    assert "=SUM(1,2)" in extract_document(tmp_path / "sample.xlsx", None, None)
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 100 Td (PDF proof) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    (tmp_path / "sample.pdf").write_bytes(buffer.getvalue())
    assert "PDF proof" in extract_document(tmp_path / "sample.pdf", 1, None)
    app = build_application(Settings(tmp_path))
    try:
        branch = app.fork()
        original = (tmp_path / "sample.pdf").read_bytes()
        app.snapshots.write_bytes(tmp_path / "sample.pdf", b"changed", app._emit(branch))
        change = next(e for e in app.events(branch) if e.kind == "file.written")
        app.restore(branch, change.id)
        assert (tmp_path / "sample.pdf").read_bytes() == original
    finally:
        app.close()


@pytest.mark.parametrize("transport", ["stdio", "streamable_http", "sse"])
def test_real_mcp_discovery_call_write_error_and_cleanup(tmp_path, transport):
    async def verify(config):
        app = build_application(Settings(tmp_path / "workspace"), mcp_servers=config)
        branch = app.fork()
        ctx = context(app, branch)
        try:
            listed = await app.runner.registry.execute("mcp_list_tools", {"server": "fixture"}, ctx)
            assert not listed.is_error, listed.content
            assert {t["name"] for t in json.loads(listed.content)["tools"]} >= {"add", "write_note", "read_reference"}
            called = await app.runner.registry.execute(
                "mcp_call_tool", {"server": "fixture", "tool": "add", "arguments": {"a": 12, "b": 25}}, ctx
            )
            assert not called.is_error and "37" in called.content
            write = await app.runner.registry.execute(
                "mcp_call_tool",
                {
                    "server": "fixture",
                    "tool": "write_note",
                    "arguments": {"filename": "external.txt", "text": "actual write"},
                },
                ctx,
            )
            assert not write.is_error and (tmp_path / "external.txt").read_text() == "actual write"
            failed = await app.runner.registry.execute(
                "mcp_call_tool", {"server": "fixture", "tool": "fail", "arguments": {}}, ctx
            )
            assert failed.is_error and "never-show-this" not in json.dumps(app.mcp_status())
            assert any(e.kind == "mcp.response" for e in app.events(branch))
            sessions = app.mcp_status()
            app.reconfigure(app.settings, app.runner.provider)
            assert app.mcp_status() == sessions
        finally:
            await app.aclose()
        assert not any(r.is_open for r in app.runner.registry.resources)

    async def scenario():
        if transport == "stdio":
            await verify(stdio_config(tmp_path))
        else:
            from sse_starlette.sse import AppStatus

            AppStatus.should_exit_event = None
            AppStatus.should_exit = False
            fixture = make_server(tmp_path)
            server_app = fixture.streamable_http_app() if transport == "streamable_http" else fixture.sse_app()
            async with LocalServer(server_app) as server:
                await verify(
                    {
                        "fixture": {
                            "transport": transport,
                            "url": server.url + ("/mcp" if transport == "streamable_http" else "/sse"),
                        }
                    }
                )

    asyncio.run(scenario())


def test_mcp_config_preserves_and_redacts_credentials(tmp_path):
    async def scenario():
        async with gateway(tmp_path) as (client, app):
            config = stdio_config(tmp_path)
            assert (await client.post("/api/mcp", json={"servers": config})).status_code == 200
            assert "never-show-this" not in (await client.get("/api/mcp")).text
            config["fixture"]["env"] = None
            assert (await client.post("/api/mcp", json={"servers": config})).status_code == 200
            private = json.loads((tmp_path / "mcp-servers.json").read_text())
            assert private["fixture"]["env"]["FIXTURE_SECRET"] == "never-show-this"
            assert (await client.post("/api/mcp/fixture/connect", json={})).status_code == 200
            assert (await client.post("/api/mcp/fixture/connect", json={"disconnect": True})).status_code == 200

    asyncio.run(scenario())


def test_user_stop_during_real_mcp_preserves_uncertain_result_and_closes(tmp_path):
    class SlowProvider:
        async def complete(self, messages, tools, on_delta):
            return ModelReply(
                Message(
                    "assistant",
                    tool_calls=(
                        ToolCall(
                            "slow-call",
                            "mcp_call_tool",
                            {"server": "fixture", "tool": "slow", "arguments": {"seconds": 30}},
                        ),
                    ),
                ),
                "tool_calls",
                Usage(10, 2),
            )

    async def scenario():
        app = build_application(Settings(tmp_path / "workspace"), SlowProvider(), stdio_config(tmp_path))
        try:
            await app.mcp_connect("fixture")
            task = asyncio.create_task(app.run("Exercise an actual slow MCP call"))
            async with asyncio.timeout(5):
                while not (tmp_path / "slow-started.txt").exists():
                    await asyncio.sleep(0.02)
            assert app.stop()
            result = await asyncio.wait_for(task, 2)
            assert result.status == "stopped" and result.model_calls == 1
            messages = [e.payload for e in app.events(result.branch_id) if e.kind == "message"]
            assert messages[-1]["role"] == "tool" and "unknown" in messages[-1]["content"]
            assert (tmp_path / "slow-started.txt").read_text() == "started"
        finally:
            await asyncio.wait_for(app.aclose(), 5)
        assert not any(r.is_open for r in app.runner.registry.resources)

    asyncio.run(scenario())


def test_mcp_config_errors_do_not_echo_secret():
    with pytest.raises(ValueError) as caught:
        normalize_servers({"broken": {"transport": "invalid", "headers": {"Authorization": "secret-value"}}})
    assert "secret-value" not in str(caught.value)


def test_real_shared_browser_controls_upload_download_assert_and_capture(tmp_path):
    async def scenario():
        app = build_application(Settings(tmp_path))
        (tmp_path / "page.html").write_text(
            """<!doctype html><title>Browser fixture</title>
        <input id="name"><span id="counter">0</span>
        <button id="add" onclick="counter.textContent=Number(counter.textContent)+1">Add</button>
        <input id="file" type="file"><input id="checked" type="checkbox">
        <select id="choice"><option value="one">One</option><option value="two">Two</option></select>
        <a id="download" download="proof.txt" href="data:text/plain,downloaded-proof">Download</a>""",
            encoding="utf-8",
        )
        branch = app.fork()
        ctx = context(app, branch)
        try:
            opened = await app.runner.registry.execute("browser_open", {"url": "page.html"}, ctx)
            assert not opened.is_error and "Browser fixture" in opened.content
            await app.browser_control({"action": "click", "target": "#add"}, branch)
            await app.browser_control({"action": "fill", "target": "#name", "value": "actual input"}, branch)
            await app.browser_control({"action": "select", "target": "#choice", "value": "two"}, branch)
            await app.browser_control({"action": "check", "target": "#checked"}, branch)
            checked = await app.browser_control({"action": "assert_text", "target": "#counter", "value": "1"}, branch)
            assert checked["assertion"]["passed"]
            uploaded = app.upload("upload.txt", b"user attachment")
            await app.browser_control({"action": "upload", "target": "#file", "files": [uploaded["path"]]}, branch)
            evaluation = await app.browser_control(
                {"action": "evaluate", "value": "document.querySelector('#file').files[0].name"}, branch
            )
            assert evaluation["evaluation"] == "upload.txt"
            download = await app.browser_control({"action": "download", "target": "#download"}, branch)
            assert (tmp_path / download["download"]["path"]).read_bytes() == b"downloaded-proof"
            screenshot = await app.runner.registry.execute("browser_action", {"action": "screenshot"}, ctx)
            picture = tmp_path / json.loads(screenshot.content)["screenshot_path"]
            assert picture.read_bytes().startswith(b"\x89PNG") and (await app.browser_image()).startswith(b"\xff\xd8")
            assert any(e.kind == "browser.manual" for e in app.events(branch))
            assert any(e.kind == "browser.action" for e in app.events(branch))
        finally:
            await app.aclose()

    asyncio.run(scenario())
