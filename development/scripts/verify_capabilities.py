"""Isolated Edge checks and real-model capability/large-task evaluations. Never patch the agent to pass."""

import argparse
import asyncio
import hashlib
import io
import json
import re
import shutil
import sys
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import FastAPI, Request
from fastapi.responses import HTMLResponse
from playwright.async_api import async_playwright, expect
from verify_frontend import BrowserTestProvider, LocalServer, send, state, wait_finished

from clearact_lite.gateway.server import create_app
from clearact_lite.settings import Settings


def credentials(workspace: Path) -> Settings:
    path = Path(__file__).resolve().parents[2] / ".clearact/web-settings.json"
    base = Settings.from_env(workspace)
    if not path.exists():
        return base
    raw = json.loads(path.read_text(encoding="utf-8"))
    fields = {
        k: raw[k]
        for k in [
            "base_url",
            "model",
            "api_key",
            "context_tokens",
            "max_output_tokens",
            "reasoning_effort",
            "temperature",
            "request_timeout_seconds",
        ]
        if k in raw
    }
    settings = replace(base, **fields)
    if urlsplit(settings.base_url).hostname not in {"localhost", "127.0.0.1", "::1"}:
        raise ValueError("This evaluation only uses the configured local model")
    return settings


def make_uploads(root: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    folder = root / "upload-sources"
    folder.mkdir()
    (folder / "orders.csv").write_text("category,amount\nA,12.5\nB,7\nA,2.5\n", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Project code: ALPHA-2026")
    doc.save(folder / "brief.docx")
    book = Workbook()
    book.active["A1"], book.active["B1"] = "adjustment", 3
    book.save(folder / "adjustment.xlsx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 100 Td (PDF proof code: PAPER-52) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    (folder / "proof.pdf").write_bytes(data.getvalue())
    return sorted(folder.iterdir())


def form_server(root: Path) -> FastAPI:
    app = FastAPI()

    @app.get("/", response_class=HTMLResponse)
    async def index():
        return HTMLResponse("""<!doctype html><meta charset="utf-8"><title>本地登记表</title>
        <style>body{font:18px system-ui;margin:40px}label{display:block;margin:20px 0}button{padding:10px}</style>
        <h1>工作登记</h1><label>金额<input id="amount" type="number"></label>
        <label>项目编号<input id="reference"></label><button id="submit">提交登记</button>
        <p id="receipt"></p><script>document.querySelector('#submit').onclick=async()=>{
        const body={amount:Number(document.querySelector('#amount').value),
        reference:document.querySelector('#reference').value};
        const result=await fetch('/receipt',{method:'POST',headers:{'Content-Type':'application/json'},
        body:JSON.stringify(body)});
        if(result.ok)document.querySelector('#receipt').textContent='登记成功：'+body.amount+' / '+body.reference;
        };</script>""")

    @app.post("/receipt")
    async def receipt(request: Request):
        data = await request.json()
        (root / "submitted-receipt.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return {"ok": True}

    @app.get("/favicon.ico")
    async def favicon():
        from fastapi.responses import Response

        return Response(status_code=204)

    return app


async def post(page, path: str, body: dict):
    return await page.evaluate(
        """async ({path,body})=>{
        const r=await fetch(path,{method:'POST',headers:{'Content-Type':'application/json',
        'X-ClearAct-Token':document.querySelector('meta[name=clearact-token]').content},body:JSON.stringify(body)});
        const data=await r.json();if(!r.ok)throw Error(data.error||'request failed');return data;
    }""",
        {"path": path, "body": body},
    )


async def report(page, root: Path, filename: str):
    current = await state(page)
    if not current["branches"]:
        return None
    branch = current["active_branch"] or current["branches"][0]["id"]
    data = await page.evaluate("id=>fetch('/api/branches/'+id).then(r=>r.json())", branch)
    (root / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


async def evaluate(mode: str, root: Path, browser):
    workspace = root / "workspace"
    settings = credentials(workspace) if mode != "ui" else Settings(workspace)
    application = create_app(
        settings, root / "web-settings.json", (lambda _: BrowserTestProvider()) if mode == "ui" else None
    )
    inputs = make_uploads(root) if mode != "large" else []
    initial_hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
    result = {
        "mode": mode,
        "provider": "protocol_test_double" if mode == "ui" else "real_local_model",
        "model": settings.model,
        "context_tokens": settings.context_tokens,
        "passed": False,
        "checks": [],
    }
    async with LocalServer(application) as gateway, LocalServer(form_server(root)) as form:
        context = await browser.new_context(viewport={"width": 1500, "height": 1000}, accept_downloads=True)
        page = await context.new_page()
        errors, console_errors = [], []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        try:
            await page.goto(gateway.url)
            await expect(page.locator("#send")).to_be_enabled()
            if mode != "large":
                # Exercise the actual editor and local stdio protocol through the UI.
                await page.locator("#mcp-trigger").click()
                await page.locator("#mcp-add").click()
                editor = page.locator("#mcp-editor")
                await editor.locator('input[name="name"]').fill("fixture")
                await editor.locator('input[name="command"]').fill(sys.executable)
                args = [str(Path(__file__).with_name("mcp_fixture.py").resolve()), "--root", str(root / "external")]
                await editor.locator('textarea[name="args"]').fill(json.dumps(args))
                await editor.locator('button[type="submit"]').click()
                await expect(editor).to_be_hidden()
                row = page.locator('.mcp-server[data-server="fixture"]')
                await row.get_by_role("button", name="连接", exact=True).click()
                await expect(row).to_contain_text("已连接", timeout=20000)
                await page.screenshot(path=str(root / "mcp-connected.png"))
                await page.locator('[data-close="mcp-dialog"]').first.click()
                result["checks"].append("MCP_editor_and_real_stdio_connection")
                await page.locator("#file-input").set_input_files([str(p) for p in inputs])
                await expect(page.locator("#attachment-list .attachment-pill")).to_have_count(4)
                await expect(page.locator("#send")).to_be_enabled()
                await page.screenshot(path=str(root / "uploads-ready.png"))
                result["checks"].append("four_format_uploads")
            if mode == "ui":
                await page.locator("#browser-toggle").click()
                await page.locator("#browser-url").fill(form.url)
                await page.locator('#browser-address-form button[type="submit"]').click()
                await expect(page.locator("#browser-image")).to_be_visible(timeout=20000)
                await expect(page.locator("#browser-feedback")).to_have_text("", timeout=10000)
                await expect(page.locator("#browser-image")).to_have_attribute("src", re.compile("^blob:"))
                await page.wait_for_function("document.querySelector('#browser-image').naturalWidth > 0")
                await page.screenshot(path=str(root / "shared-browser.png"))
                result["checks"].append("browser_live_preview")
                # Click the visible screenshot at the real page's first input, then type through the UI.
                info = await page.evaluate("fetch('/api/browser/state').then(r=>r.json())")
                controls = await post(
                    page,
                    "/api/browser/action",
                    {
                        "arguments": {
                            "action": "evaluate",
                            "value": "(()=>{const r=document.querySelector('#amount').getBoundingClientRect();"
                            "return {x:r.x+r.width/2,y:r.y+r.height/2}})()",
                        }
                    },
                )
                point = controls["evaluation"]
                image = page.locator("#browser-image")
                rect = await image.bounding_box()
                await image.click(
                    position={
                        "x": point["x"] / info["viewport"]["width"] * rect["width"],
                        "y": point["y"] / info["viewport"]["height"] * rect["height"],
                    }
                )
                await expect(page.locator("#browser-feedback")).to_have_text("", timeout=10000)
                await page.locator("#browser-text").fill("62")
                await page.locator('#browser-type-form button[type="submit"]').click()
                await expect(page.locator("#browser-feedback")).to_have_text("", timeout=10000)
                value = await post(
                    page,
                    "/api/browser/action",
                    {"arguments": {"action": "evaluate", "value": "document.querySelector('#amount').value"}},
                )
                assert value["evaluation"] == "62"
                result["checks"].append("user_screenshot_click_and_keyboard_share_session")
                await page.locator("#browser-panel-close").click()
                await send(page, "记录本轮附件并创建一个交付文件")
                await wait_finished(page)
                saved = await report(page, root, "ui-report.json")
                assert any(e["kind"] == "attachment.linked" for e in saved["events"])
                result["checks"].append("attachment_message_and_delivery")
                mobile = await browser.new_context(viewport={"width": 390, "height": 844}, is_mobile=True)
                mobile_page = await mobile.new_page()
                await mobile_page.goto(gateway.url)
                await expect(mobile_page.locator("#send")).to_be_enabled()
                assert await mobile_page.evaluate("document.documentElement.scrollWidth <= innerWidth")
                await mobile_page.screenshot(path=str(root / "mobile-capabilities.png"))
                await mobile.close()
                result["checks"].append("mobile_controls_fit")
                result["passed"] = True
            else:
                if mode == "large":
                    goal = (Path(__file__).resolve().parents[1] / "fixtures/large-visual-novel.txt").read_text(
                        encoding="utf-8"
                    )
                    result["goal_sha256"] = hashlib.sha256(goal.encode()).hexdigest()
                    (root / "goal.txt").write_text(goal, encoding="utf-8")
                    print("LARGE_TASK: original visual-novel goal, empty workspace", flush=True)
                else:
                    goal = (
                        "读取本轮四个附件：CSV 中汇总 amount，Excel 读取 adjustment，"
                        "DOCX 读取项目编号，PDF 读取校验码。"
                        "使用 MCP 服务 fixture 的 read_reference 获得 amount；"
                        "用该服务 add 计算 CSV合计+adjustment+MCP金额。"
                        f"然后用浏览器打开 {form.url}，在金额框填总金额，项目编号填 DOCX 编号，"
                        "点击提交登记，实际核对登记成功。"
                        "生成 report.json，键精确为 csv_total、adjustment、mcp_amount、grand_total、"
                        "docx_code、pdf_code。"
                        "保存一张浏览器截图，与 report.json 一同登记交付。保持附件不变，完成后简短报告。"
                    )
                    print("REAL_CAPABILITIES: uploads -> MCP -> browser form -> files", flush=True)
                await send(page, goal)
                await wait_finished(page, 360000 if mode == "large" else 240000)
                saved = await report(page, root, "report.json")
                result["runs"] = saved["runs"]
                result["tools"] = [e["payload"]["name"] for e in saved["events"] if e["kind"] == "tool.started"]
                assert saved["runs"][-1]["status"] == "finished", saved["runs"][-1]
                if mode == "real":
                    actual = json.loads((workspace / "report.json").read_text(encoding="utf-8"))
                    assert actual == {
                        "csv_total": 22,
                        "adjustment": 3,
                        "mcp_amount": 37,
                        "grand_total": 62,
                        "docx_code": "ALPHA-2026",
                        "pdf_code": "PAPER-52",
                    }, actual
                    assert json.loads((root / "submitted-receipt.json").read_text()) == {
                        "amount": 62,
                        "reference": "ALPHA-2026",
                    }
                    assert "mcp_call_tool" in result["tools"] and "browser_action" in result["tools"]
                    assert result["tools"].count("document_read") >= 3
                    links = next(e for e in saved["events"] if e["kind"] == "attachment.linked")["payload"]["files"]
                    assert all(
                        hashlib.sha256((workspace / f["path"]).read_bytes()).hexdigest() == initial_hashes[f["name"]]
                        for f in links
                    )
                    assert len(saved["artifacts"]) >= 2
                    result["checks"].extend(
                        [
                            "native_document_and_CSV_reading",
                            "real_MCP_computation",
                            "real_browser_form_submission",
                            "independent_report_values",
                            "original_uploads_unchanged",
                            "screenshot_and_report_delivery",
                        ]
                    )
                    await page.locator("#browser-toggle").click()
                    await expect(page.locator("#browser-image")).to_be_visible(timeout=15000)
                    await page.screenshot(path=str(root / "real-shared-browser.png"))
                    result["passed"] = True
                else:
                    result["generation_finished"] = True
                    result["independent_acceptance"] = "pending"
                    result["passed"] = None
            result["page_errors"], result["console_errors"] = errors, console_errors
            assert not errors
        except BaseException as exc:
            result["error"] = f"{type(exc).__name__}: {exc}"
            try:
                await post(page, "/api/stop", {})
                saved = await report(page, root, "failure-report.json")
                if saved:
                    result["runs"] = saved["runs"]
            except Exception:
                pass
            await page.screenshot(path=str(root / "failure.png"))
            raise
        finally:
            (root / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
            await context.close()


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["ui", "real", "large"], default="ui")
    parser.add_argument("--feedback-case", type=Path)
    args = parser.parse_args()
    root = (
        Path(__file__).resolve().parents[1]
        / "records/evals"
        / datetime.now().strftime(f"capability-{'feedback' if args.feedback_case else args.mode}-%Y%m%d-%H%M%S")
    )
    root.mkdir(parents=True)
    print("RECORDS:", root, flush=True)
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(channel="msedge", headless=True)
        try:
            if args.feedback_case:
                await evaluate_feedback(args.feedback_case.resolve(), root, browser)
            else:
                await evaluate(args.mode, root, browser)
        finally:
            await browser.close()
    print("VERIFICATION:", "explicit_feedback" if args.feedback_case else args.mode, "finished", flush=True)


async def evaluate_feedback(original: Path, root: Path, browser):
    workspace = original / "workspace"
    baseline = json.loads((original / "report.json").read_text(encoding="utf-8"))
    branch = baseline["branch_id"]
    shutil.copy2(original / "mcp-servers.json", root / "mcp-servers.json")
    shutil.copy2(workspace / "report.json", root / "before-report.json")
    application = create_app(credentials(workspace), root / "web-settings.json")
    result = {
        "mode": "explicit_feedback",
        "original_case": str(original),
        "passed": False,
        "initial_attempt_passed": False,
        "automatic_runtime_repair": False,
    }
    async with LocalServer(application) as gateway, LocalServer(form_server(root)) as form:
        context = await browser.new_context(viewport={"width": 1500, "height": 1000})
        page = await context.new_page()
        try:
            await page.goto(gateway.url)
            await expect(page.locator("#send")).to_be_enabled()
            await expect(page.locator("#topic-count")).to_have_text("1")
            prompt = (
                "独立核验指出上一轮算错：12.5+7+2.5 应为22，不是22.5。"
                "请用 fixture MCP 的 add 逐项计算 CSV 合计，再加 Excel 的3与 MCP金额37，得到正确总额。"
                f"登记页面测试服务已重启，当前地址是 {form.url}。"
                "在此页面重新实际提交正确金额和原项目编号，检查成功结果。"
                "修正 report.json，重新截图并登记交付两个文件，保持附件不变。"
            )
            (root / "feedback.txt").write_text(prompt, encoding="utf-8")
            await send(page, prompt)
            await wait_finished(page, 240000)
            data = await report(page, root, "report.json")
            result["runs"] = data["runs"]
            assert data["branch_id"] == branch and data["runs"][-1]["status"] == "finished"
            actual = json.loads((workspace / "report.json").read_text(encoding="utf-8"))
            assert actual == {
                "csv_total": 22,
                "adjustment": 3,
                "mcp_amount": 37,
                "grand_total": 62,
                "docx_code": "ALPHA-2026",
                "pdf_code": "PAPER-52",
            }
            assert json.loads((root / "submitted-receipt.json").read_text()) == {
                "amount": 62,
                "reference": "ALPHA-2026",
            }
            files = next(e for e in data["events"] if e["kind"] == "attachment.linked")["payload"]["files"]
            assert all(hashlib.sha256((workspace / f["path"]).read_bytes()).hexdigest() == f["sha256"] for f in files)
            await page.locator("#browser-toggle").click()
            await expect(page.locator("#browser-image")).to_be_visible(timeout=15000)
            await page.wait_for_function("document.querySelector('#browser-image').naturalWidth > 0")
            await page.screenshot(path=str(root / "corrected-browser.png"))
            result["passed"] = True
        except BaseException as exc:
            result["error"] = f"{type(exc).__name__}: {exc}"
            await post(page, "/api/stop", {})
            await report(page, root, "failure-report.json")
            await page.screenshot(path=str(root / "failure.png"))
            raise
        finally:
            (root / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
            await context.close()


if __name__ == "__main__":
    asyncio.run(main())
